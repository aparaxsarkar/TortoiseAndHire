"""`ExportService` round-trip against a real database (ADR-0005 revised, ADR-0010,
ADR-0013 - Title/Company/URL corrections)."""

from __future__ import annotations

import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.exports.excel import read_workbook
from app.models import Application, ApplicationImportRun, Company, Job, Source, SourcePosting
from app.schemas.applications import ApplicationPatch
from app.services.applications import ApplicationService
from app.services.exports import ExportService
from tests.integration.conftest import SessionFactory


def _seed(session: Session) -> dict[str, uuid.UUID]:
    company = Company(name="Acme", normalized_name=f"acme-{uuid.uuid4()}")
    session.add(company)
    session.flush()
    a = Job(company_id=company.id, title="ML Engineer", normalized_title="ml engineer")
    b = Job(company_id=company.id, title="Backend Engineer", normalized_title="backend engineer")
    session.add_all([a, b])
    session.flush()
    return {"a": a.id, "b": b.id, "company": company.id}


def _add_posting(
    session: Session, job_id: uuid.UUID, *, slug: str, canonical_url: str, dedup_key: str
) -> SourcePosting:
    source = Source(slug=f"{slug}-{uuid.uuid4()}", display_name=slug.title())
    session.add(source)
    session.flush()
    posting = SourcePosting(
        source_id=source.id,
        job_id=job_id,
        source_job_id=dedup_key,
        canonical_url=canonical_url,
        dedup_key=dedup_key,
        raw_payload={},
        content_hash="h",
    )
    session.add(posting)
    session.flush()
    return posting


def _edit_cell(data: bytes, *, row: int, header: str, value: object) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Applications"]
    col = [c.value for c in next(ws.iter_rows(max_row=1))].index(header) + 1
    ws.cell(row=row, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row_of(data: bytes, job_id: uuid.UUID) -> int:
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Applications"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    jid_col = headers.index("job_id") + 1
    return next(
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=jid_col).value == str(job_id)
    )


def test_export_snapshots_every_job(db_session: Session, session_factory: SessionFactory) -> None:
    ids = _seed(db_session)
    ApplicationService(session_factory=session_factory).patch(
        ids["a"], ApplicationPatch(applied=True, outcome="screen")
    )

    filename, data = ExportService(session_factory=session_factory).export_xlsx()
    assert filename.endswith(".xlsx")

    rows = {r.job_id: r for r in read_workbook(data)}
    assert set(rows) == {str(ids["a"]), str(ids["b"])}
    assert rows[str(ids["a"])].values["applied"] is True
    assert rows[str(ids["a"])].values["outcome"] == "screen"
    assert rows[str(ids["a"])].revision == 1
    assert rows[str(ids["b"])].revision == 0  # no application row yet


def test_dry_run_reports_changes_without_writing(
    db_session: Session, session_factory: SessionFactory
) -> None:
    _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(data, row=2, header="Outcome", value="onsite")

    report = exports.import_xlsx(filename="edit.xlsx", data=edited, commit=False)

    assert report.committed is False
    by_action = {r.action for r in report.rows}
    assert "updated" in by_action or "created" in by_action
    changed = next(r for r in report.rows if r.action in ("updated", "created"))
    assert any(c.field == "outcome" and c.new == "onsite" for c in changed.changes)

    # nothing hit the database
    assert db_session.execute(select(func.count()).select_from(Application)).scalar_one() == 0
    assert (
        db_session.execute(select(func.count()).select_from(ApplicationImportRun)).scalar_one() == 0
    )


def test_commit_applies_changes_and_writes_an_audit_row(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    # rows are ordered by company then title: "Backend Engineer" (b) is row 2
    edited = _edit_cell(data, row=2, header="Applied", value="yes")
    edited = _edit_cell(edited, row=2, header="Notes", value="applied via referral")

    report = exports.import_xlsx(filename="edit.xlsx", data=edited, commit=True)

    assert report.committed is True
    assert report.counts["created"] == 1

    app = db_session.execute(select(Application).where(Application.job_id == ids["b"])).scalar_one()
    assert app.applied is True
    assert app.notes == "applied via referral"
    assert app.updated_via == "excel_import"
    assert app.revision == 1

    audit = db_session.execute(select(ApplicationImportRun)).scalar_one()
    assert audit.committed is True
    assert audit.stats["created"] == 1


def test_a_stale_revision_is_a_conflict_and_is_skipped(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    apps = ApplicationService(session_factory=session_factory)
    exports = ExportService(session_factory=session_factory)
    apps.patch(ids["a"], ApplicationPatch(outcome="screen"))  # revision 1

    _, data = exports.export_xlsx()  # the sheet records revision 1 for job a
    apps.patch(ids["a"], ApplicationPatch(notes="a later edit"))  # DB revision -> 2

    edited = _edit_cell(data, row=_row_of(data, ids["a"]), header="Outcome", value="onsite")
    report = exports.import_xlsx(filename="stale.xlsx", data=edited, commit=True)

    conflict = next(r for r in report.rows if r.job_id == ids["a"])
    assert conflict.action == "conflict"

    app = db_session.execute(select(Application).where(Application.job_id == ids["a"])).scalar_one()
    assert app.outcome == "screen"  # untouched by the import
    assert app.notes == "a later edit"
    assert app.revision == 2


def test_unknown_job_id_and_bad_enum_are_error_rows(
    db_session: Session, session_factory: SessionFactory
) -> None:
    _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()

    data = _edit_cell(data, row=2, header="job_id", value=str(uuid.uuid4()))  # unknown job
    data = _edit_cell(data, row=3, header="Networking", value="mystery")  # bad enum

    report = exports.import_xlsx(filename="bad.xlsx", data=data, commit=True)

    errors = [r for r in report.rows if r.action == "error"]
    assert len(errors) == 2
    assert any("no such job" in (e.message or "") for e in errors)
    assert any("networking" in (e.message or "") for e in errors)


# --- ADR-0013: canonical field corrections (Title / Company / URL) ---------


def test_title_only_edit_is_reported_separately_and_not_applied_on_dry_run(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(
        data, row=_row_of(data, ids["b"]), header="Title", value="Backend Engineer II"
    )

    report = exports.import_xlsx(filename="title.xlsx", data=edited, commit=False)

    row = next(r for r in report.rows if r.job_id == ids["b"])
    assert row.action == "unchanged"  # no application-field change
    assert row.changes == []
    assert [c.field for c in row.canonical_changes] == ["title"]
    assert row.canonical_changes[0].old == "Backend Engineer"
    assert row.canonical_changes[0].new == "Backend Engineer II"
    assert report.canonical_change_count == 1
    assert db_session.get(Job, ids["b"]).title == "Backend Engineer"  # dry run: untouched


def test_title_edit_is_applied_on_commit_and_renormalized(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(
        data, row=_row_of(data, ids["b"]), header="Title", value="Backend Engineer II"
    )

    report = exports.import_xlsx(filename="title.xlsx", data=edited, commit=True)

    job = db_session.get(Job, ids["b"])
    assert job.title == "Backend Engineer II"
    assert job.normalized_title == "backend engineer ii"
    # no Application row was created - only a canonical field changed
    assert db_session.execute(select(func.count()).select_from(Application)).scalar_one() == 0
    assert next(r for r in report.rows if r.job_id == ids["b"]).action == "unchanged"


def test_company_edit_reassigns_the_job_rather_than_renaming(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(data, row=_row_of(data, ids["a"]), header="Company", value="Acme AI")

    report = exports.import_xlsx(filename="company.xlsx", data=edited, commit=True)

    row = next(r for r in report.rows if r.job_id == ids["a"])
    assert [c.field for c in row.canonical_changes] == ["company_name"]

    job_a = db_session.get(Job, ids["a"])
    job_b = db_session.get(Job, ids["b"])
    original_company = db_session.get(Company, ids["company"])

    assert job_a.company_id != ids["company"]  # reassigned to a new company
    new_company = db_session.get(Company, job_a.company_id)
    assert new_company.name == "Acme AI"
    assert original_company.name == "Acme"  # the shared row was never renamed
    assert job_b.company_id == ids["company"]  # the sibling job is untouched


def test_url_edit_applies_when_exactly_one_posting_and_never_touches_dedup_key(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    posting = _add_posting(
        db_session,
        ids["a"],
        slug="greenhouse",
        canonical_url="https://boards.greenhouse.io/acme/1",
        dedup_key="G-1",
    )
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(
        data,
        row=_row_of(data, ids["a"]),
        header="URL",
        value="https://boards.greenhouse.io/acme/999",
    )

    report = exports.import_xlsx(filename="url.xlsx", data=edited, commit=True)

    row = next(r for r in report.rows if r.job_id == ids["a"])
    assert [c.field for c in row.canonical_changes] == ["url"]

    db_session.refresh(posting)
    assert posting.canonical_url == "https://boards.greenhouse.io/acme/999"
    assert posting.dedup_key == "G-1"  # identity is untouched (ADR-0002)


def test_url_edit_is_ignored_when_the_job_has_no_posting(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)  # job "b" has zero linked postings
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(
        data, row=_row_of(data, ids["b"]), header="URL", value="https://example.com/job/1"
    )

    report = exports.import_xlsx(filename="url.xlsx", data=edited, commit=True)

    row = next(r for r in report.rows if r.job_id == ids["b"])
    assert row.canonical_changes == []
    assert row.message is not None and "0 linked posting" in row.message


def test_url_edit_is_ignored_when_the_job_has_multiple_postings(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    p1 = _add_posting(
        db_session,
        ids["a"],
        slug="greenhouse",
        canonical_url="https://boards.greenhouse.io/acme/1",
        dedup_key="G-1",
    )
    p2 = _add_posting(
        db_session,
        ids["a"],
        slug="lever",
        canonical_url="https://jobs.lever.co/acme/1",
        dedup_key="L-1",
    )
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(
        data, row=_row_of(data, ids["a"]), header="URL", value="https://example.com/somewhere-else"
    )

    report = exports.import_xlsx(filename="url.xlsx", data=edited, commit=True)

    row = next(r for r in report.rows if r.job_id == ids["a"])
    assert row.canonical_changes == []
    assert row.message is not None and "2 linked posting" in row.message

    db_session.refresh(p1)
    db_session.refresh(p2)
    assert p1.canonical_url == "https://boards.greenhouse.io/acme/1"
    assert p2.canonical_url == "https://jobs.lever.co/acme/1"


def test_an_invalid_url_is_an_error_row(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    edited = _edit_cell(data, row=_row_of(data, ids["a"]), header="URL", value="not a url")

    report = exports.import_xlsx(filename="bad-url.xlsx", data=edited, commit=True)

    # a row-level parse_error short-circuits before job_id is resolved, same as
    # any other malformed cell (e.g. a bad "Applied" value) - so find it by action.
    errors = [r for r in report.rows if r.action == "error"]
    assert len(errors) == 1
    assert errors[0].job_id is None
    assert "URL" in (errors[0].message or "")

    # nothing was applied - not even the untouched job's fields
    assert db_session.get(Job, ids["a"]).title == "ML Engineer"


def test_application_and_canonical_edits_in_one_row_both_land(
    db_session: Session, session_factory: SessionFactory
) -> None:
    ids = _seed(db_session)
    exports = ExportService(session_factory=session_factory)
    _, data = exports.export_xlsx()
    row_num = _row_of(data, ids["b"])
    edited = _edit_cell(data, row=row_num, header="Outcome", value="screen")
    edited = _edit_cell(edited, row=row_num, header="Title", value="Backend Engineer II")

    report = exports.import_xlsx(filename="both.xlsx", data=edited, commit=True)

    row = next(r for r in report.rows if r.job_id == ids["b"])
    assert row.action == "created"  # the application-field change created the row
    assert [c.field for c in row.changes] == ["outcome"]
    assert [c.field for c in row.canonical_changes] == ["title"]
    assert report.canonical_change_count == 1

    app = db_session.execute(select(Application).where(Application.job_id == ids["b"])).scalar_one()
    assert app.outcome == "screen"
    assert db_session.get(Job, ids["b"]).title == "Backend Engineer II"
