"""`GET /exports/xlsx` + `POST /exports/import` - protected (ADR-0005 revised, ADR-0010)."""

from __future__ import annotations

import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_export_service
from app.exports.excel import read_workbook
from app.main import create_app
from app.models import Application, Company, Job
from app.services.exports import ExportService
from tests.integration.conftest import SessionFactory

TOKEN = "s3cr3t-test-token"
AUTH = {"Authorization": f"Bearer {TOKEN}"}
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _token_env(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("TORTOISEANDHIRE_API_TOKEN", TOKEN)


@pytest.fixture
def client(session_factory: SessionFactory) -> TestClient:
    app = create_app()
    app.dependency_overrides[get_export_service] = lambda: ExportService(
        session_factory=session_factory
    )
    return TestClient(app)


def _seed(session: Session) -> uuid.UUID:
    company = Company(name="Acme", normalized_name=f"acme-{uuid.uuid4()}")
    session.add(company)
    session.flush()
    job = Job(company_id=company.id, title="ML Engineer", normalized_title="ml engineer")
    session.add(job)
    session.flush()
    return job.id


def test_download_requires_the_token(client: TestClient) -> None:
    assert client.get("/api/v1/exports/xlsx").status_code == 401


def test_download_returns_a_spreadsheet(client: TestClient, db_session: Session) -> None:
    _seed(db_session)
    resp = client.get("/api/v1/exports/xlsx", headers=AUTH)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == _XLSX
    assert "attachment" in resp.headers["content-disposition"]
    assert len(read_workbook(resp.content)) == 1


def test_import_empty_body_is_a_422_problem(client: TestClient) -> None:
    resp = client.post("/api/v1/exports/import", headers=AUTH, content=b"")
    assert resp.status_code == 422
    assert resp.headers["content-type"].startswith("application/problem+json")


def test_reimporting_an_untouched_sheet_is_all_unchanged(
    client: TestClient, db_session: Session
) -> None:
    _seed(db_session)
    xlsx = client.get("/api/v1/exports/xlsx", headers=AUTH).content

    resp = client.post("/api/v1/exports/import?commit=false", headers=AUTH, content=xlsx)
    assert resp.status_code == 200
    body = resp.json()
    assert body["committed"] is False
    assert body["counts"]["unchanged"] == 1
    assert body["counts"]["updated"] == 0


def test_commit_writes_the_change_through(client: TestClient, db_session: Session) -> None:
    job_id = _seed(db_session)
    xlsx = client.get("/api/v1/exports/xlsx", headers=AUTH).content

    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb["Applications"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    ws.cell(row=2, column=headers.index("Applied") + 1, value="yes")
    ws.cell(row=2, column=headers.index("Notes") + 1, value="from the spreadsheet")
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post("/api/v1/exports/import?commit=true", headers=AUTH, content=buf.getvalue())
    assert resp.status_code == 200
    assert resp.json()["counts"]["created"] == 1

    app = db_session.execute(select(Application).where(Application.job_id == job_id)).scalar_one()
    assert app.applied is True
    assert app.notes == "from the spreadsheet"
    assert app.updated_via == "excel_import"


def test_garbage_upload_is_a_422_problem(client: TestClient) -> None:
    resp = client.post("/api/v1/exports/import", headers=AUTH, content=b"nonsense")
    assert resp.status_code == 422
    assert resp.headers["content-type"].startswith("application/problem+json")
