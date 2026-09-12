"""`.xlsx` <-> rows, openpyxl only. No DB, no business rules.

`write_workbook` renders `ExportRow`s. `read_workbook` parses an uploaded file
back into `ParsedRow`s (hidden keys + lightly-coerced editable cells); the
service does the DB reconcile and enum validation.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import TortoiseError
from app.exports.layout import COLUMNS, EDITABLE_COLUMNS, HEADERS
from app.schemas.exports import ExportRow

_SHEET = "Applications"
_TRUE = {"true", "yes", "y", "1", "x", "applied"}
_FALSE = {"false", "no", "n", "0", "", "not applied"}


class WorkbookError(TortoiseError):
    http_status = 422
    http_title = "Unprocessable Entity"


@dataclass
class ParsedRow:
    excel_row: int  # 1-based
    job_id: str | None
    revision: int | None
    values: dict[str, Any] = field(default_factory=dict)  # editable field -> coerced value
    parse_error: str | None = None  # a bad cell in this row; the service reports it


def _cell_out(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)  # openpyxl stores naive datetimes
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


def write_workbook(rows: list[ExportRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert isinstance(ws, Worksheet)  # a fresh Workbook always has one worksheet
    ws.title = _SHEET
    ws.append(list(HEADERS))
    for idx, col in enumerate(COLUMNS, start=1):
        if col.hidden:
            ws.column_dimensions[get_column_letter(idx)].hidden = True
    for row in rows:
        ws.append([_cell_out(getattr(row, col.field)) for col in COLUMNS])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _as_bool(value: Any) -> bool | None:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in _TRUE:
        return True
    if text in _FALSE:
        return False
    raise WorkbookError(f"{value!r} is not a yes/no value")


def _as_datetime(value: Any) -> dt.datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value if value.tzinfo else value.replace(tzinfo=dt.UTC)
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day, tzinfo=dt.UTC)
    try:
        return dt.datetime.fromisoformat(str(value)).replace(tzinfo=dt.UTC)
    except ValueError as exc:
        raise WorkbookError(f"{value!r} is not a date/time") from exc


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _as_applied(value: Any) -> bool:
    # a blank cell means "not applied"; genuine garbage still raises
    return bool(_as_bool(value))


def _as_enum_text(value: Any) -> str:
    return _as_text(value) or "none"


def _as_url(value: Any) -> str | None:
    # blank means "leave as-is" (canonical_url is NOT NULL); garbage is rejected
    # outright rather than silently corrupting a working link.
    text = _as_text(value)
    if text is None:
        return None
    parsed = urlsplit(text)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise WorkbookError(f"{text!r} is not a valid http(s) URL")
    return text


_COERCE: dict[str, Callable[[Any], Any]] = {
    "title": _as_text,
    "company_name": _as_text,
    "url": _as_url,
    "applied": _as_applied,
    "applied_at": _as_datetime,
    "networking": _as_enum_text,
    "outcome": _as_enum_text,
    "application_url": _as_text,
    "notes": _as_text,
}


def read_workbook(data: bytes) -> list[ParsedRow]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises a grab-bag of types for bad files
        raise WorkbookError(f"could not open the workbook: {exc}") from exc

    ws = wb[_SHEET] if _SHEET in wb.sheetnames else wb.active
    if not isinstance(ws, Worksheet):
        raise WorkbookError("no readable worksheet in the workbook")

    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        raise WorkbookError("the workbook is empty") from None

    pos = {name: i for i, name in enumerate(header)}
    for required in ("job_id",):
        if required not in pos:
            raise WorkbookError(f"missing the {required!r} column - export a fresh sheet")

    def at(cells: tuple[Any, ...], name: str) -> Any:
        i = pos.get(name)
        return cells[i] if i is not None and i < len(cells) else None

    out: list[ParsedRow] = []
    for excel_row, cells in enumerate(rows, start=2):
        if cells is None or all(c is None for c in cells):
            continue
        raw_job_id = at(cells, "job_id")
        raw_rev = at(cells, "revision")
        values: dict[str, Any] = {}
        parse_error: str | None = None
        for col in EDITABLE_COLUMNS:
            try:
                values[col.field] = _COERCE[col.field](at(cells, col.header))
            except WorkbookError as exc:
                parse_error = f"{col.header}: {exc}"
                break
        out.append(
            ParsedRow(
                excel_row=excel_row,
                job_id=(str(raw_job_id).strip() or None) if raw_job_id is not None else None,
                revision=int(raw_rev) if isinstance(raw_rev, (int, float)) else None,
                values=values,
                parse_error=parse_error,
            )
        )
    return out
